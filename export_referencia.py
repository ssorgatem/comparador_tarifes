"""
Genera un JSON de referència amb els valors ja calculats del full "Facturas"
d'un Excel real (amb macros habilitades, ja recalculat), perquè el test de
validació completa (`tests/test_validacio_excel.py`) els pugui fer servir.

Aquest JSON conté les TEVES dades de consum real -> NO el pugis a un
repositori públic. Es genera sempre en local, a `tests/fixtures/`, que està
exclòs a `.gitignore`.

Ús:
    python3 tools/export_referencia.py RUTA_AL_XLSM_JA_CALCULAT
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl

COL_LETTERS = ["E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P"]


def export(xlsm_path: str) -> dict:
    wb = openpyxl.load_workbook(xlsm_path, keep_vba=True, data_only=True)
    wf = wb["Facturas"]

    referencia = {}
    for row in range(4, wf.max_row + 1):
        comercializadora = wf.cell(row=row, column=1).value
        tarifa = wf.cell(row=row, column=2).value
        if not comercializadora and not tarifa:
            continue
        valors = []
        for col in COL_LETTERS:
            v = wf[f"{col}{row}"].value
            valors.append(None if v in (None, "", " ") else float(v))
        referencia[str(row)] = valors

    return referencia


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsm_path")
    ap.add_argument("--out", default=str(Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "facturas_referencia.json"))
    args = ap.parse_args()

    referencia = export(args.xlsm_path)

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(referencia, f, ensure_ascii=False, indent=2)

    print(f"Exportades {len(referencia)} files de Facturas -> {out_path}", file=sys.stderr)
    print("Recorda: també cal tests/fixtures/tarifas.json i tests/fixtures/datos.json "
          "(genera'ls amb extract_tarifas.py / extract_datos.py sobre el mateix Excel).", file=sys.stderr)


if __name__ == "__main__":
    main()
