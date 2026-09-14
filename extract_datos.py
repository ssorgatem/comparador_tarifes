"""
Extreu el full "Datos" del Simulador d'Excel cap a un JSON estructurat:
constants globals (alquiler, bono social, IE, IVA, potència FV) + una llista
de "mesos" amb els consums de cada columna (D, E, F, ... una per mes).

Ús:
    python3 extract_datos.py RUTA_AL_XLSM [--out datos.json]
"""
import argparse
import json
import sys

import openpyxl
from openpyxl.utils import get_column_letter

FIRST_MONTH_COL = 4  # columna D

# fila -> camp, per a cada columna de mes
MONTH_ROWS = {
    9: "dias",
    10: "pot_p1",
    11: "pot_p2",
    12: "consumo_p1",
    13: "consumo_p2",
    14: "consumo_p3",
    16: "kwh_exc",
    17: "kwh_barato_nocturna",
    18: "kwh_caro_nocturna",
    20: "kwh_barato_ve",
    21: "kwh_caro_ve",
    23: "kwh_barato_solar",
    24: "kwh_caro_solar",
    26: "kwh_50h_gratis",
    29: "kwh_2h_gratis",
    32: "kwh_3h_gratis",
    35: "kwh_promo_sunclub",
    36: "kwh_gratis_sunclub",
    37: "kwh_no_promo_sunclub",
    5: "autoconsumo_estimado",
}

GLOBAL_CELLS = {
    "alquiler_contador_dia": "B1",
    "bono_social_dia": "B2",
    "potencia_fotovoltaica_kw": "B3",
    "impuesto_electrico_pct": "B4",
    "iva_pct": "B5",
}


def extract(xlsm_path: str) -> dict:
    wb = openpyxl.load_workbook(xlsm_path, keep_vba=True, data_only=True)
    ws = wb["Datos"]

    globals_ = {name: ws[coord].value for name, coord in GLOBAL_CELLS.items()}

    meses = []
    col = FIRST_MONTH_COL
    while True:
        # La fila 7 numera els mesos (1, 2, 3...) fins la columna "SUMA"; és la
        # marca de fi de taula més fiable (Descripción i Días poden estar buits
        # si encara no s'hi ha enganxat cap dada, com en una plantilla verge).
        num_mes = ws.cell(row=7, column=col).value
        if not isinstance(num_mes, (int, float)):
            break
        desc = ws.cell(row=8, column=col).value  # "Descripción" a la fila 8 (D8, E8, ...)
        mes = {"columna": get_column_letter(col), "num_mes": num_mes, "descripcion": desc}
        for row, field in MONTH_ROWS.items():
            mes[field] = ws.cell(row=row, column=col).value
        meses.append(mes)
        col += 1

    return {"globals": globals_, "meses": meses}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsm_path")
    ap.add_argument("--out", default="datos.json")
    args = ap.parse_args()

    data = extract(args.xlsm_path)

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"Globals: {data['globals']}", file=sys.stderr)
    print(f"Extrets {len(data['meses'])} mesos -> {args.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
