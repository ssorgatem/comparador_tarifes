"""
Extreu el full "Tarifas" del Simulador d'Excel cap a un JSON estructurat.

Es pot re-executar contra qualsevol versió futura del fitxer .xlsm: només cal
que el full es continuï dient "Tarifas" i mantingui les mateixes columnes
(A..R, capçalera a la fila 4, dades a partir de la fila 5).

Ús:
    python3 extract_tarifas.py RUTA_AL_XLSM [--out tarifas.json]
"""
import argparse
import json
import sys

import openpyxl

# Columna (1-indexada) -> nom de camp, seguint exactament l'ordre en què
# FACTURA() els consumeix a Módulo1.bas
COLUMNS = {
    1: "comercializadora",
    2: "tarifa",
    3: "precio_potencia_punta",       # €/kW y año
    4: "precio_potencia_valle",       # €/kW y año
    5: "precio_energia_punta",        # c€/kWh
    6: "precio_energia_llana",        # c€/kWh
    7: "precio_energia_valle",        # c€/kWh
    8: "precio_excedente",            # c€/kWh
    9: "metodo",                      # dispatch: None/"PVPC"/"Consumo" -> genèric; "BateriaXX" -> específic
    10: "franjas_horarias",           # "Normal" / "Nocturna-14h" / "VE-6h" / "Solar-16h" / ...
    11: "extra_mensual",              # €/mes (gestión)
    12: "extra_kwp",                  # €/kWp
    13: "extra_excedentes",           # €
    14: "info_tecnica_indexada",      # "campo churro": IDX=...|DSV=...|FNEE=...
    15: "descuento_factura",          # €
    16: "cif",
    17: "enlace",
    18: "mas_info",
}

FIRST_DATA_ROW = 5


def extract(xlsm_path: str) -> list[dict]:
    wb = openpyxl.load_workbook(xlsm_path, keep_vba=True, data_only=True)
    ws = wb["Tarifas"]

    tarifas = []
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, max_row=ws.max_row):
        comercializadora = row[0].value
        tarifa = row[1].value
        if not comercializadora and not tarifa:
            continue  # fila buida (format sense dades)

        entry = {}
        for cell in row:
            field = COLUMNS.get(cell.column)
            if field is None:
                continue
            entry[field] = cell.value
        entry["_fila_excel"] = row[0].row  # útil per depurar / retrobar la fila original
        tarifas.append(entry)

    return tarifas


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsm_path")
    ap.add_argument("--out", default="tarifas.json")
    args = ap.parse_args()

    tarifas = extract(args.xlsm_path)

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(tarifas, f, ensure_ascii=False, indent=2)

    print(f"Extretes {len(tarifas)} tarifes -> {args.out}", file=sys.stderr)

    # Petit resum de cobertura per mètode, útil per saber què falta implementar
    from collections import Counter
    c = Counter(t.get("metodo") or "(genèric)" for t in tarifas)
    print("Distribució per mètode:", file=sys.stderr)
    for metodo, n in c.most_common():
        print(f"  {metodo:20s} {n:4d}  ({100*n/len(tarifas):.0f}%)", file=sys.stderr)


if __name__ == "__main__":
    main()
